#!/usr/bin/env python3
"""Offline end-to-end test: download -> cache -> restart -> backtest -> Excel, using a fake
Breeze client. Writes only to a temporary folder; your real data/ and output/ are untouched.

Run from the project root:  python tests/smoke_test.py
"""
import os
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def run(cmd, env):
    print("$", " ".join(cmd))
    r = subprocess.run(cmd, cwd=ROOT, env=env, capture_output=True, text=True)
    if r.returncode:
        print(r.stdout, r.stderr)
        sys.exit(f"FAILED: {' '.join(cmd)}")
    return r.stdout


def check_time_exit_skips_muhurat():
    from datetime import date
    sys.path.insert(0, str(ROOT))
    from ic import config as C
    from ic.common import load_calendar, time_exit_day
    cal = load_calendar(date(2021, 10, 1), date(2021, 11, 30))
    assert date(2021, 11, 4) not in cal, "Muhurat day should not be a trading day"
    saved, C.EXIT_DTE = C.EXIT_DTE, 18
    try:   # expiry 22-Nov-2021 - 18 days = Muhurat Thu 4-Nov-2021 -> must roll back to Wed 3-Nov
        assert time_exit_day(date(2021, 11, 22), cal) == date(2021, 11, 3)
    finally:
        C.EXIT_DTE = saved
    print("time exit skips Muhurat session: OK")


def main():
    check_time_exit_skips_muhurat()
    tmp = Path(tempfile.mkdtemp(prefix="ic_smoke_"))
    env = dict(os.environ,
               PYTHONPATH=os.pathsep.join([str(ROOT / "tests" / "fake_breeze"), str(ROOT)]),
               IC_DATA_DIR=str(tmp / "data"), IC_OUTPUT_DIR=str(tmp / "output"), IC_PAUSE_SEC="0",
               BREEZE_API_KEY="test", BREEZE_API_SECRET="test", BREEZE_SESSION="test")
    out1 = run([sys.executable, "download.py", "--trades", "1", "9", "23", "62", "67"], env)
    assert "All requested data downloaded" in out1, out1
    out2 = run([sys.executable, "download.py", "--trades", "1", "9", "23", "62", "67"], env)
    assert "0 API calls needed" in out2, "restart should skip cached data"
    out4 = run([sys.executable, "download.py", "--index"], env)
    assert (tmp / "data" / "nifty_daily.csv").exists(), out4
    out3 = run([sys.executable, "backtest.py", "--trades", "1", "9", "23", "62", "67"], env)
    report = tmp / "output" / "IC_backtest_report.xlsx"
    assert report.exists(), "report not written"
    import openpyxl
    names = openpyxl.load_workbook(report).sheetnames
    assert names == ["Summary", "Trades", "Daily_MTM", "Reconciliation", "Skipped", "Data_Quality"], names
    assert "manual exit" in out3, "trade 23 should be flagged as a manual StockMock exit"
    tr = openpyxl.load_workbook(report)["Trades"]
    hdr = [c.value for c in tr[1]]
    notes = {r[hdr.index("trade_id")]: r[hdr.index("strike_note")] for r in tr.iter_rows(min_row=2, values_only=True)}
    if os.environ.get("IC_HEDGE_WIDTH", "300") == "300":   # the fake 13500PE hedge has no data
        assert notes.get(1) and "PE 13800/13500 ->" in notes[1], f"trade 1 should use fallback PE strikes: {notes}"
    out7 = run([sys.executable, "export_prices.py", "--trades", "1", "9", "23", "62", "67",
                "--out", str(tmp / "output" / "smoke_prices.xlsx")], env)
    assert (tmp / "output" / "smoke_prices.xlsx").exists(), out7
    out8 = run([sys.executable, "backtest.py", "--trades", "1", "9", "23", "62", "67",
                "--out", str(tmp / "output" / "smoke_gapup.xlsx")], dict(env, IC_GAP_UP_EXIT="1.0"))
    summary_text = [str(c.value) for row in openpyxl.load_workbook(tmp / "output" / "smoke_gapup.xlsx")["Summary"].iter_rows()
                    for c in row if c.value is not None]
    assert any("Gap-up exit: ON" in t for t in summary_text), "gap-up exit note missing from the report"
    env9 = dict(env, IC_ENTRY_SHIFT="2")
    run([sys.executable, "download.py", "--trades", "9", "62"], env9)
    out9 = run([sys.executable, "backtest.py", "--trades", "9", "62", "--out", str(tmp / "output" / "smoke_wed.xlsx")], env9)
    assert "Entry shifted 2" in "".join(str(c.value) for r in openpyxl.load_workbook(tmp / "output" / "smoke_wed.xlsx")["Summary"].iter_rows() for c in r), out9
    out5 = run([sys.executable, "select_strikes.py", "--ce-delta", "0.25", "--name", "smoke", "--trades", "9", "62"], env)
    assert (tmp / "data" / "strike_sets" / "strikes_smoke.csv").exists(), out5
    env5 = dict(env, IC_STRIKE_SET="smoke")
    run([sys.executable, "download.py", "--trades", "9", "62"], env5)
    out6 = run([sys.executable, "backtest.py", "--trades", "9", "62", "--out", str(tmp / "output" / "smoke_ce25.xlsx")], env5)
    assert (tmp / "output" / "smoke_ce25.xlsx").exists(), out6
    print(out3)
    print(f"SMOKE TEST PASSED  (temp files in {tmp})")


if __name__ == "__main__":
    main()
